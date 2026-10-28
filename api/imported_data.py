import csv
import io
import re
import unicodedata
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any



REGISTRY_FILES = (
    "Distribuidores a Granel de GLP.csv",
    "Distribuidores en Cilindros de GLP.csv",
    "Medios de Transporte Acuático.csv",
    "Transporte de combustibles líquidos en contenedores intermedios.csv",
    "Transporte GLP a Granel.csv",
    "Transporte GLP en Cilindros.csv",
    "Transporte Terrestre de Petróleo Crudo.csv",
    "Transportistas de Combustibles Líquidos y OPDH en camiones.csv",
    "Transportistas de Combustibles Líquidos y OPDH en vagones.csv",
    "CDCL y OPDH con Instalaciones Fijas_0.csv",
    "CDCL y OPDH con Instalaciones Móviles_0.csv",
    "Consumidores Menores.csv",
    "EESS con GLP y GNV_0.csv",
    "EESS con GNV_0.csv",
    "Estaciones de Servicio con Gasocentro de GLP.csv",
    "15Gasocentros_1.csv",
    "Grifos Flotantes.csv",
    "Grifos Rurales con almacenamiento en Cilindros.csv",
    "Comercializadores de Combustible de Aviación.csv",
    "Comercializadores de Combustible de Embarcación.csv",
    "Consumidor Directo de GLP.csv",
    "Distribuidores Mayoristas Combustibles Líquidos.csv",
    "Distribuidores Mayoristas OPDH.csv",
    "Grifos y Estaciones de Servicios.csv",
    "Importador de GLP.csv",
    "Locales de Venta de GLP.csv",
    "Planta de Abastecimiento GLP.csv",
    "Planta Envasadora GLP.csv",
    "Redes de Distribución de GLP.csv",
)
ENVIRONMENTAL_FILES = {
    "IPASH_Agua_0.xlsx": "Agua",
    "IPASH_Aire.xlsx": "Aire",
    "IPASH_Sedimento_1.xlsx": "Sedimento",
    "IPASH_Suelo_0.xlsx": "Suelo",
}
PRICE_FILE = "precios_combustibles_anonimizados_20260301_part1 (2).csv"
MINORIST_PRICE_FILE = "CL-Registro-precios-DMIN.csv"
PRICE_COLUMNS = {
    "g_premium": "Gasohol premium",
    "g_regular": "Gasohol regular",
    "diesel": "Diesel",
    "gnv": "GNV",
    "glp_g": "GLP granel",
    "glp_e": "GLP envasado",
}


def _normalize(value: Any) -> str:
    decomposed = unicodedata.normalize("NFD", str(value or ""))
    return "".join(char for char in decomposed if unicodedata.category(char) != "Mn").strip().upper()


def _key(value: Any) -> str:
    return re.sub(r"[^A-Z0-9]", "", _normalize(value))


def _csv_rows(path: Path):
    with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as source:
        sample = source.read(8192)
        source.seek(0)
        try:
            delimiter = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
        except csv.Error:
            first_line = next((line for line in sample.splitlines() if line.strip()), "")
            delimiter = max((",", ";", "\t", "|"), key=first_line.count)
        yield from csv.DictReader(source, delimiter=delimiter)


def _field(row: dict[str, Any], *names: str) -> Any:
    values = {_key(name): value for name, value in row.items()}
    for name in names:
        value = values.get(_key(name))
        if value is not None and str(value).strip():
            return value
    return None


def _number(value: Any) -> float | None:
    if value is None:
        return None
    cleaned = str(value).replace(",", "").strip()
    try:
        number = float(cleaned)
    except ValueError:
        return None
    return number if number == number and abs(number) != float("inf") else None


def _localized_number(value: Any) -> float | None:
    if isinstance(value, str) and "," in value and "." not in value:
        value = value.replace(",", ".")
    return _number(value)


def _date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value or "").strip()
    for format_string in ("%Y-%m-%d", "%Y%m%d", "%d/%m/%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(text[:10], format_string).date()
        except ValueError:
            continue
    return None


def _build_registry_summary(root: Path) -> dict[str, Any]:
    sources = []
    departments: Counter[str] = Counter()
    categories: Counter[str] = Counter()
    total_records = 0
    for filename in REGISTRY_FILES:
        path = root / filename
        if not path.is_file():
            continue
        record_count = 0
        source_departments: Counter[str] = Counter()
        source_categories: Counter[str] = Counter()
        for row in _csv_rows(path):
            if not any(value and str(value).strip() for value in row.values()):
                continue
            record_count += 1
            department = str(_field(row, "DEPARTAMENTO") or "").strip().title()
            category = str(_field(row, "TIPO TRANSPORTE", "ACTIVIDAD", "TIPO") or path.stem.split("_")[0]).strip()
            if department:
                source_departments[department] += 1
                departments[department] += 1
            if category:
                source_categories[category] += 1
                categories[category] += 1
        total_records += record_count
        sources.append({
            "fileName": filename,
            "records": record_count,
            "categories": [{"name": name, "records": count} for name, count in source_categories.most_common(10)],
        })
    return {
        "records": total_records,
        "sourceCount": len(sources),
        "sources": sources,
        "byDepartment": [{"name": name, "records": count} for name, count in departments.most_common(25)],
        "byCategory": [{"name": name, "records": count} for name, count in categories.most_common(25)],
        "description": "Registros de agentes, instalaciones y medios autorizados; no son movimientos de inventario ni envíos realizados.",
    }


def _build_environmental_summary(root: Path) -> dict[str, Any]:
    from openpyxl import load_workbook

    components: Counter[str] = Counter()
    parameters: Counter[tuple[str, str]] = Counter()
    years: Counter[int] = Counter()
    earliest: date | None = None
    latest: date | None = None
    total_records = 0
    sources = []

    for filename, component in ENVIRONMENTAL_FILES.items():
        path = root / filename
        if not path.is_file():
            continue
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            worksheet = workbook.active
            rows = worksheet.iter_rows(values_only=True)
            header = next(rows, ())
            indices = {_key(name): index for index, name in enumerate(header) if name is not None}
            parameter_index = indices.get("PARAMETRO")
            unit_index = indices.get("UNIDADDEMEDIDA")
            date_index = next((indices[name] for name in ("FECHA", "FECHAINICIO") if name in indices), None)
            file_count = 0
            for row in rows:
                if parameter_index is None or parameter_index >= len(row):
                    continue
                parameter = str(row[parameter_index] or "").strip()
                if not parameter:
                    continue
                unit = str(row[unit_index] or "").strip() if unit_index is not None and unit_index < len(row) else ""
                measured_at = _date(row[date_index]) if date_index is not None and date_index < len(row) else None
                file_count += 1
                total_records += 1
                components[component] += 1
                parameters[(parameter, unit)] += 1
                if measured_at:
                    years[measured_at.year] += 1
                    earliest = min(earliest, measured_at) if earliest else measured_at
                    latest = max(latest, measured_at) if latest else measured_at
            sources.append({"fileName": filename, "component": component, "records": file_count})
        finally:
            workbook.close()

    return {
        "records": total_records,
        "sourceCount": len(sources),
        "sources": sources,
        "byComponent": [{"name": name, "records": count} for name, count in components.most_common()],
        "byParameter": [
            {"name": name, "unit": unit, "records": count}
            for (name, unit), count in parameters.most_common(60)
        ],
        "byYear": [{"year": year, "records": count} for year, count in sorted(years.items())],
        "dateRange": {
            "from": earliest.isoformat() if earliest else None,
            "to": latest.isoformat() if latest else None,
        },
        "thresholdsConfigured": False,
        "description": "Muestreos ambientales históricos de pasivos; no son telemetría en tiempo real ni alertas de ruta.",
    }


def _build_price_summary(root: Path) -> dict[str, Any]:
    path = root / PRICE_FILE
    if not path.is_file():
        return {"available": False, "records": 0, "sourceFile": PRICE_FILE}

    per_product: dict[str, dict[str, dict[str, float | int]]] = defaultdict(dict)
    total_records = 0
    earliest: date | None = None
    latest: date | None = None
    for row in _csv_rows(path):
        total_records += 1
        evaluated_at = _date(_field(row, "fe_eval"))
        if evaluated_at:
            earliest = min(earliest, evaluated_at) if earliest else evaluated_at
            latest = max(latest, evaluated_at) if latest else evaluated_at
        if not evaluated_at:
            continue
        period = evaluated_at.isoformat()
        for column, label in PRICE_COLUMNS.items():
            value = _number(_field(row, column))
            if value is None:
                continue
            stats = per_product[label].setdefault(period, {"count": 0, "sum": 0.0, "minimum": value, "maximum": value})
            stats["count"] = int(stats["count"]) + 1
            stats["sum"] = float(stats["sum"]) + value
            stats["minimum"] = min(float(stats["minimum"]), value)
            stats["maximum"] = max(float(stats["maximum"]), value)

    latest_prices = []
    for product, periods in sorted(per_product.items()):
        if not periods:
            continue
        period = max(periods)
        stats = periods[period]
        count = int(stats["count"])
        latest_prices.append({
            "product": product,
            "period": period,
            "observations": count,
            "average": float(stats["sum"]) / count,
            "minimum": stats["minimum"],
            "maximum": stats["maximum"],
        })
    minorist_path = root / MINORIST_PRICE_FILE
    minorist_products: dict[str, dict[str, float | int]] = {}
    minorist_earliest: date | None = None
    minorist_latest: date | None = None
    if minorist_path.is_file():
        for row in _csv_rows(minorist_path):
            product = str(_field(row, "PRODUCTO") or "").strip()
            minimum = _localized_number(_field(row, "PRECIO_MIN_SOLES"))
            maximum = _localized_number(_field(row, "PRECIO_MAX_SOLES"))
            if not product or minimum is None or maximum is None:
                continue
            stats = minorist_products.setdefault(product, {
                "count": 0, "midpointSum": 0.0, "minimum": minimum, "maximum": maximum,
            })
            stats["count"] = int(stats["count"]) + 1
            stats["midpointSum"] = float(stats["midpointSum"]) + (minimum + maximum) / 2
            stats["minimum"] = min(float(stats["minimum"]), minimum)
            stats["maximum"] = max(float(stats["maximum"]), maximum)
            registered_at = _date(_field(row, "FECHA_DE_REGISTRO"))
            if registered_at:
                minorist_earliest = min(minorist_earliest, registered_at) if minorist_earliest else registered_at
                minorist_latest = max(minorist_latest, registered_at) if minorist_latest else registered_at
    return {
        "available": True,
        "sourceFile": PRICE_FILE,
        "records": total_records,
        "dateRange": {
            "from": earliest.isoformat() if earliest else None,
            "to": latest.isoformat() if latest else None,
        },
        "latestPrices": latest_prices,
        "minoristRegistry": {
            "sourceFile": MINORIST_PRICE_FILE,
            "records": sum(int(stats["count"]) for stats in minorist_products.values()),
            "dateRange": {
                "from": minorist_earliest.isoformat() if minorist_earliest else None,
                "to": minorist_latest.isoformat() if minorist_latest else None,
            },
            "byProduct": [
                {
                    "product": product,
                    "records": int(stats["count"]),
                    "averageBandMidpoint": float(stats["midpointSum"]) / int(stats["count"]),
                    "minimum": stats["minimum"],
                    "maximum": stats["maximum"],
                    "unit": "Galones",
                }
                for product, stats in sorted(minorist_products.items())
            ],
        },
        "description": "Precios anonimizados; la unidad requiere validar el diccionario de la fuente y no representan inventario ni transacciones identificadas.",
    }


def build_imported_data(root: Path) -> dict[str, dict[str, Any]]:
    return {
        "registryData": _build_registry_summary(root),
        "environmentalData": _build_environmental_summary(root),
        "priceObservations": _build_price_summary(root),
    }

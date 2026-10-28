import csv
import hashlib
import io
import json
import logging
import os
import re
import secrets
import shutil
import time
import unicodedata
import mimetypes
from contextlib import contextmanager
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from typing import Any, Callable, Iterator

import psycopg
from psycopg.rows import dict_row
from psycopg.types.json import Jsonb
from api.imported_data import (
    ENVIRONMENTAL_FILES,
    MINORIST_PRICE_FILE,
    PRICE_FILE,
    REGISTRY_FILES,
    build_imported_data,
)
from api.operational_data import build_operational_data

IMPORTS_PATH = Path(__file__).resolve().parents[1] / "public" / "imports"
REPORTS_PATH = Path(os.environ.get("REPORTS_PATH", Path(__file__).resolve().parents[1] / "reportes_generados" / "files"))
SESSION_TTL = timedelta(hours=8)
logger = logging.getLogger(__name__)

DATASET_FILES = {
    "facilities": "CDCLyOPDHconInstalacionesFijas.csv",
    "complaints": "DenunciasSINADA20192025_2944631782350634676.csv",
    "fuelDictionary": "Diccionario_de_datos_Precios_Combustible.csv",
    "cases": "Expedientesresueltos20212025_2944571782350726004.csv",
    "resolutions": "Resolucionesdirectorales20182025_2944601782350740761.csv",
}
ALTERNATE_COMPLAINTS_FILE = "DenunciasSINADA_2898361792458095761.csv"
OPERATIONAL_PREFIXES = (
    "MNGN2-3-", "MNGN4-", "MNGN5-", "MNGN6-", "MNGN7-", "MNGN8-", "MNHL2-",
    "MNHL3-", "MNHL4-", "MNHL6-", "MNHL7-", "MNHL8-", "MNHL9-",
)

PERMISSIONS_BY_ROLE = {
    "Gerente de Operaciones": ["dashboard", "trazabilidad", "reportes", "fuentes"],
    "Analista de Datos": ["dashboard", "riesgos", "trazabilidad", "reportes", "fuentes"],
    "Supervisor de Cumplimiento": [
        "dashboard", "riesgos", "auditoria", "trazabilidad",
        "cumplimiento", "reportes", "supervision", "fuentes",
    ],
    "Ingeniero de Campo": ["dashboard", "riesgos", "trazabilidad", "reportes", "fuentes"],
}

REPORT_TYPES_BY_ROLE = {
    "Gerente de Operaciones": [
        "Dashboard Ejecutivo", "Reporte de Riesgos", "Reporte de Trazabilidad", "Reporte Ejecutivo",
    ],
    "Analista de Datos": [
        "Reporte de Riesgos", "Reporte de Trazabilidad", "Reporte de Inventario",
        "Reporte de Transporte", "Reporte de Predicción de Demanda", "Reportes Analíticos",
    ],
    "Supervisor de Cumplimiento": [
        "Dashboard Ejecutivo", "Reporte de Cumplimiento", "Reporte de Riesgos",
        "Reporte de Auditoría", "Reporte de Trazabilidad", "Evidencias Digitales",
    ],
    "Ingeniero de Campo": [
        "Reporte de Transporte", "Reporte de Riesgos de Transporte",
        "Reporte de Condiciones Ambientales", "Reporte de Alertas", "Reporte de Trazabilidad",
    ],
}


@contextmanager
def connection() -> Iterator[psycopg.Connection[dict[str, Any]]]:
    database_url = os.environ.get("DATABASE_URL", "").strip()
    if not database_url:
        raise RuntimeError("Configura DATABASE_URL para conectar con PostgreSQL.")
    with psycopg.connect(database_url, row_factory=dict_row) as conn:
        yield conn


def require_database_url() -> None:
    if not os.environ.get("DATABASE_URL", "").strip():
        raise RuntimeError("Configura DATABASE_URL para conectar con PostgreSQL.")


def wait_for_connection(max_attempts: int = 30) -> None:
    for attempt in range(max_attempts):
        try:
            with connection() as conn:
                conn.execute("SELECT 1")
            return
        except psycopg.OperationalError:
            if attempt == max_attempts - 1:
                raise
            logger.warning("PostgreSQL aún no está disponible; reintentando (%s/%s).", attempt + 1, max_attempts)
            time.sleep(1)


def initialize_database(on_authentication_ready: Callable[[], None] | None = None) -> None:
    startup_started = time.perf_counter()
    logger.info("Esperando conexión con PostgreSQL.")
    wait_for_connection()
    connection_ready = time.perf_counter()
    logger.info("PostgreSQL conectado; preparando esquema.")
    with connection() as conn:
        statements = (
            """
            CREATE TABLE IF NOT EXISTS sector_datasets (
                dataset_key TEXT PRIMARY KEY,
                source_files JSONB NOT NULL,
                checksum TEXT NOT NULL,
                rows JSONB NOT NULL,
                updated_at TIMESTAMPTZ NOT NULL DEFAULT now()
            );
            CREATE TABLE IF NOT EXISTS managed_users (
                username TEXT PRIMARY KEY,
                display_name TEXT NOT NULL,
                password_hash TEXT NOT NULL,
                role TEXT NOT NULL,
                permissions JSONB NOT NULL,
                active BOOLEAN NOT NULL DEFAULT TRUE,
                created_at TIMESTAMPTZ NOT NULL DEFAULT now()
            )
            """,
            """
            CREATE TABLE IF NOT EXISTS auth_sessions (
                token_hash TEXT PRIMARY KEY,
                username TEXT NOT NULL REFERENCES managed_users(username) ON DELETE CASCADE,
                expires_at TIMESTAMPTZ NOT NULL,
                created_at TIMESTAMPTZ NOT NULL DEFAULT now()
            )
            """,
            "CREATE INDEX IF NOT EXISTS auth_sessions_expires_at_idx ON auth_sessions(expires_at)",
            """
            CREATE TABLE IF NOT EXISTS access_events (
                id BIGSERIAL PRIMARY KEY,
                occurred_at TIMESTAMPTZ NOT NULL DEFAULT now(),
                username TEXT NOT NULL,
                role TEXT NOT NULL,
                action TEXT NOT NULL,
                resource TEXT NOT NULL,
                result TEXT NOT NULL CHECK (result IN ('Permitido', 'Denegado'))
            )
            """,
            "CREATE INDEX IF NOT EXISTS access_events_occurred_at_idx ON access_events(occurred_at DESC)",
            """
            CREATE TABLE IF NOT EXISTS source_files (
                file_name TEXT PRIMARY KEY,
                file_type TEXT NOT NULL,
                checksum TEXT NOT NULL,
                size_bytes BIGINT NOT NULL,
                modified_time_ns BIGINT,
                content BYTEA,
                updated_at TIMESTAMPTZ NOT NULL DEFAULT now()
            )
            """,
            "ALTER TABLE source_files ADD COLUMN IF NOT EXISTS modified_time_ns BIGINT",
            "ALTER TABLE source_files ALTER COLUMN content DROP NOT NULL",
            """
            CREATE TABLE IF NOT EXISTS generated_reports (
                id TEXT PRIMARY KEY,
                file_name TEXT NOT NULL UNIQUE,
                name TEXT NOT NULL,
                report_type TEXT NOT NULL,
                period TEXT NOT NULL,
                username TEXT NOT NULL,
                role TEXT NOT NULL,
                sections JSONB NOT NULL,
                size_bytes BIGINT NOT NULL,
                created_at TIMESTAMPTZ NOT NULL DEFAULT now()
            )
            """,
            """
            CREATE TABLE IF NOT EXISTS risk_assessments (
                id TEXT PRIMARY KEY,
                scenario_name TEXT NOT NULL,
                inputs JSONB NOT NULL,
                score INTEGER NOT NULL CHECK (score BETWEEN 0 AND 100),
                risk_level TEXT NOT NULL CHECK (risk_level IN ('Bajo', 'Medio', 'Alto', 'Crítico')),
                factors JSONB NOT NULL,
                assessment_type TEXT NOT NULL CHECK (assessment_type = 'simulación'),
                created_by TEXT NOT NULL,
                created_at TIMESTAMPTZ NOT NULL DEFAULT now()
            )
            """,
            "CREATE INDEX IF NOT EXISTS risk_assessments_created_at_idx ON risk_assessments(created_at DESC)",
            "ALTER TABLE generated_reports DROP CONSTRAINT IF EXISTS generated_reports_username_fkey",
            """
            CREATE TABLE IF NOT EXISTS ai_governance_controls (
                control_id TEXT PRIMARY KEY,
                standard TEXT NOT NULL,
                model_scope TEXT NOT NULL DEFAULT 'ambos',
                focus TEXT NOT NULL,
                objective TEXT NOT NULL,
                status TEXT NOT NULL DEFAULT 'no_evaluado'
                    CHECK (status IN ('no_evaluado', 'en_revision', 'conforme', 'brecha')),
                evidence TEXT NOT NULL DEFAULT '',
                updated_by TEXT,
                updated_at TIMESTAMPTZ NOT NULL DEFAULT now(),
                CHECK (status <> 'conforme' OR length(trim(evidence)) > 0)
            )
            """,
            "ALTER TABLE ai_governance_controls ADD COLUMN IF NOT EXISTS model_scope TEXT NOT NULL DEFAULT 'ambos'",
            """
            CREATE TABLE IF NOT EXISTS ai_governance_evidence_files (
                id TEXT PRIMARY KEY,
                control_id TEXT NOT NULL REFERENCES ai_governance_controls(control_id) ON DELETE CASCADE,
                file_name TEXT NOT NULL,
                media_type TEXT NOT NULL,
                size_bytes INTEGER NOT NULL CHECK (size_bytes BETWEEN 1 AND 10485760),
                checksum TEXT NOT NULL,
                content BYTEA NOT NULL,
                uploaded_by TEXT NOT NULL,
                created_at TIMESTAMPTZ NOT NULL DEFAULT now()
            )
            """,
            "CREATE INDEX IF NOT EXISTS ai_governance_evidence_control_idx ON ai_governance_evidence_files(control_id, created_at DESC)",
            """
            CREATE TABLE IF NOT EXISTS ai_governance_events (
                id BIGSERIAL PRIMARY KEY,
                control_id TEXT NOT NULL REFERENCES ai_governance_controls(control_id),
                username TEXT NOT NULL,
                previous_status TEXT NOT NULL,
                new_status TEXT NOT NULL,
                evidence TEXT NOT NULL,
                occurred_at TIMESTAMPTZ NOT NULL DEFAULT now()
            )
            """,
            "CREATE INDEX IF NOT EXISTS ai_governance_events_control_idx ON ai_governance_events(control_id, occurred_at DESC)",
            """
            CREATE TABLE IF NOT EXISTS ai_governance_risks (
                id TEXT PRIMARY KEY,
                model_name TEXT NOT NULL,
                category TEXT NOT NULL,
                description TEXT NOT NULL,
                data_basis TEXT NOT NULL CHECK (data_basis IN ('observado', 'sintetico', 'conceptual')),
                evidence TEXT NOT NULL DEFAULT '',
                likelihood SMALLINT NOT NULL CHECK (likelihood BETWEEN 1 AND 5),
                impact SMALLINT NOT NULL CHECK (impact BETWEEN 1 AND 5),
                owner TEXT NOT NULL,
                treatment TEXT NOT NULL,
                status TEXT NOT NULL DEFAULT 'identificado'
                    CHECK (status IN ('identificado', 'en_tratamiento', 'aceptado', 'cerrado')),
                created_by TEXT NOT NULL,
                created_at TIMESTAMPTZ NOT NULL DEFAULT now(),
                updated_at TIMESTAMPTZ NOT NULL DEFAULT now()
            )
            """,
            "CREATE INDEX IF NOT EXISTS ai_governance_risks_updated_at_idx ON ai_governance_risks(updated_at DESC)",
            """
            CREATE TABLE IF NOT EXISTS ai_governance_risk_events (
                id BIGSERIAL PRIMARY KEY,
                risk_id TEXT NOT NULL REFERENCES ai_governance_risks(id) ON DELETE CASCADE,
                username TEXT NOT NULL,
                previous_status TEXT,
                new_status TEXT NOT NULL,
                treatment TEXT NOT NULL,
                occurred_at TIMESTAMPTZ NOT NULL DEFAULT now()
            )
            """,
        )
        for statement in statements:
            conn.execute(statement)

    initialize_demand_model_schema()
    initialize_ai_governance_schema()
    with connection() as conn:
        conn.execute(
            "UPDATE managed_users SET permissions = permissions || %s WHERE NOT (permissions ? %s)",
            (Jsonb(["fuentes"]), "fuentes"),
        )
    seed_initial_supervisor()
    logger.info("Esquema y usuario inicial listos; autenticación disponible.")
    if on_authentication_ready is not None:
        on_authentication_ready()

    REPORTS_PATH.mkdir(parents=True, exist_ok=True)
    source_sync_started = time.perf_counter()
    logger.info("Iniciando sincronización del catálogo de fuentes.")
    changed_sources, source_checksums = import_source_files()
    logger.info(
        "Catálogo de fuentes sincronizado; %s/%s archivos nuevos o modificados. Importando datasets.",
        len(changed_sources),
        len(source_checksums),
    )
    import_source_datasets(changed_sources, source_checksums)
    logger.info(
        "Fuentes sincronizadas: %s/%s cambiaron; PostgreSQL listo en %.2f s, sincronización en %.2f s.",
        len(changed_sources),
        len(source_checksums),
        connection_ready - startup_started,
        time.perf_counter() - source_sync_started,
    )
    migrate_legacy_reports()


def initialize_ai_governance_schema() -> None:
    controls = (
        ("27001-access", "ISO/IEC 27001:2022", "ambos", "Seguridad de la información", "Definir y revisar el acceso a datos, modelos y funciones según responsabilidades."),
        ("27001-trace", "ISO/IEC 27001:2022", "ambos", "Registro y trazabilidad", "Conservar evidencia de versiones, fuentes, evaluaciones y cambios de los modelos."),
        ("27001-monitor", "ISO/IEC 27001:2022", "ambos", "Monitoreo e incidentes", "Definir cómo detectar, registrar y atender incidentes de seguridad o fallas del sistema de IA."),
        ("38507-accountability", "ISO/IEC 38507:2022", "ambos", "Responsabilidad de gobernanza", "Asignar responsables organizacionales para aprobar, supervisar y revisar el uso de IA."),
        ("38507-human-oversight", "ISO/IEC 38507:2022", "ambos", "Supervisión y decisión humana", "Establecer cuándo una recomendación requiere revisión humana y quién puede aceptarla o descartarla."),
        ("38507-risk-impact", "ISO/IEC 38507:2022", "ambos", "Riesgo e impacto", "Evaluar impactos, límites de uso, incertidumbre y riesgos del sistema durante su ciclo de vida."),
        ("9001-quality-objectives", "ISO 9001:2026", "demanda", "Calidad predictiva de demanda", "Registrar MAE, RMSE, MAPE/WAPE y R² en una prueba temporal y compararlos con el baseline."),
        ("9001-road-quality", "ISO 9001:2026", "vial", "Calidad y alcance del modelo vial", "Documentar la métrica frente al baseline y que el objetivo es clasificar emergencias reportadas, no estimar riesgo por viaje."),
        ("9001-improvement", "ISO 9001:2026", "ambos", "No conformidad y mejora", "Registrar desviaciones, acciones correctivas, responsables y revisión de su eficacia."),
    )
    with connection() as conn:
        with conn.cursor() as cursor:
            cursor.executemany(
                """
                INSERT INTO ai_governance_controls (control_id, standard, model_scope, focus, objective)
                VALUES (%s, %s, %s, %s, %s)
                ON CONFLICT (control_id) DO UPDATE SET
                    standard = EXCLUDED.standard,
                    model_scope = EXCLUDED.model_scope,
                    focus = EXCLUDED.focus,
                    objective = EXCLUDED.objective
                """,
                controls,
            )


def list_ai_governance_controls(scope: str = "todos") -> list[dict[str, Any]]:
    with connection() as conn:
        return conn.execute(
            """
            SELECT control_id AS id, standard, model_scope, focus, objective, status, evidence,
                   updated_by, updated_at
            FROM ai_governance_controls
            WHERE %s = 'todos' OR model_scope IN (%s, 'ambos')
            ORDER BY standard, control_id
            """,
            (scope, scope),
        ).fetchall()


def update_ai_governance_control(
    control_id: str,
    status: str,
    evidence: str,
    username: str,
) -> dict[str, Any] | None:
    with connection() as conn:
        current = conn.execute(
            "SELECT status FROM ai_governance_controls WHERE control_id = %s FOR UPDATE",
            (control_id,),
        ).fetchone()
        if not current:
            return None
        updated = conn.execute(
            """
            UPDATE ai_governance_controls
            SET status = %s, evidence = %s, updated_by = %s, updated_at = now()
            WHERE control_id = %s
            RETURNING control_id AS id, standard, model_scope, focus, objective, status, evidence,
                      updated_by, updated_at
            """,
            (status, evidence, username, control_id),
        ).fetchone()
        conn.execute(
            """
            INSERT INTO ai_governance_events
                (control_id, username, previous_status, new_status, evidence)
            VALUES (%s, %s, %s, %s, %s)
            """,
            (control_id, username, current["status"], status, evidence),
        )
        return updated


def list_ai_governance_evidence_files(control_id: str | None = None) -> list[dict[str, Any]]:
    with connection() as conn:
        if control_id:
            return conn.execute(
                """
                SELECT id, control_id, file_name, media_type, size_bytes, checksum,
                       uploaded_by, created_at
                FROM ai_governance_evidence_files
                WHERE control_id = %s
                ORDER BY created_at DESC, id
                """,
                (control_id,),
            ).fetchall()
        return conn.execute(
            """
            SELECT id, control_id, file_name, media_type, size_bytes, checksum,
                   uploaded_by, created_at
            FROM ai_governance_evidence_files
            ORDER BY control_id, created_at DESC, id
            """
        ).fetchall()


def save_ai_governance_evidence_file(
    *,
    evidence_id: str,
    control_id: str,
    file_name: str,
    media_type: str,
    content: bytes,
    checksum: str,
    username: str,
) -> dict[str, Any] | None:
    with connection() as conn:
        control = conn.execute(
            "SELECT status, evidence FROM ai_governance_controls WHERE control_id = %s FOR UPDATE",
            (control_id,),
        ).fetchone()
        if not control:
            return None
        attachment_reference = f"Adjunto: {file_name} · SHA-256 {checksum}"
        evidence = control["evidence"].strip()
        evidence_changed = attachment_reference not in evidence
        if evidence_changed:
            evidence = "\n".join(part for part in (evidence, attachment_reference) if part)
        if len(evidence) > 5000:
            raise ValueError("La evidencia supera el límite de 5000 caracteres al añadir el archivo.")
        row = conn.execute(
            """
            INSERT INTO ai_governance_evidence_files
                (id, control_id, file_name, media_type, size_bytes, checksum, content, uploaded_by)
            VALUES (%s, %s, %s, %s, %s, %s, %s, %s)
            RETURNING id, control_id, file_name, media_type, size_bytes, checksum,
                      uploaded_by, created_at
            """,
            (evidence_id, control_id, file_name, media_type, len(content), checksum, content, username),
        ).fetchone()
        if evidence_changed:
            conn.execute(
                "UPDATE ai_governance_controls SET evidence = %s, updated_by = %s, updated_at = now() WHERE control_id = %s",
                (evidence, username, control_id),
            )
        conn.execute(
            """
            INSERT INTO ai_governance_events
                (control_id, username, previous_status, new_status, evidence)
            VALUES (%s, %s, %s, %s, %s)
            """,
            (control_id, username, control["status"], control["status"], attachment_reference),
        )
        return {**row, "control_evidence": evidence}


def get_ai_governance_evidence_file(evidence_id: str) -> dict[str, Any] | None:
    with connection() as conn:
        return conn.execute(
            """
            SELECT id, control_id, file_name, media_type, size_bytes, checksum,
                   content, uploaded_by, created_at
            FROM ai_governance_evidence_files WHERE id = %s
            """,
            (evidence_id,),
        ).fetchone()


def list_ai_governance_risks() -> list[dict[str, Any]]:
    with connection() as conn:
        return conn.execute(
            """
            SELECT id, model_name, category, description, data_basis, evidence,
                   likelihood, impact, owner, treatment, status, created_by,
                   created_at, updated_at
            FROM ai_governance_risks
            ORDER BY CASE status WHEN 'identificado' THEN 0 WHEN 'en_tratamiento' THEN 1 ELSE 2 END,
                     impact * likelihood DESC, updated_at DESC
            """
        ).fetchall()


def create_ai_governance_risk(risk: dict[str, Any], username: str) -> dict[str, Any]:
    with connection() as conn:
        row = conn.execute(
            """
            INSERT INTO ai_governance_risks
                (id, model_name, category, description, data_basis, evidence,
                 likelihood, impact, owner, treatment, created_by)
            VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
            RETURNING id, model_name, category, description, data_basis, evidence,
                      likelihood, impact, owner, treatment, status, created_by,
                      created_at, updated_at
            """,
            (
                risk["id"], risk["modelName"], risk["category"], risk["description"],
                risk["dataBasis"], risk["evidence"], risk["likelihood"], risk["impact"],
                risk["owner"], risk["treatment"], username,
            ),
        ).fetchone()
        conn.execute(
            "INSERT INTO ai_governance_risk_events (risk_id, username, new_status, treatment) VALUES (%s, %s, %s, %s)",
            (risk["id"], username, "identificado", risk["treatment"]),
        )
        return row


def update_ai_governance_risk(
    risk_id: str, status: str, treatment: str, username: str
) -> dict[str, Any] | None:
    with connection() as conn:
        current = conn.execute(
            "SELECT status FROM ai_governance_risks WHERE id = %s FOR UPDATE", (risk_id,)
        ).fetchone()
        if not current:
            return None
        row = conn.execute(
            """
            UPDATE ai_governance_risks SET status = %s, treatment = %s, updated_at = now()
            WHERE id = %s
            RETURNING id, model_name, category, description, data_basis, evidence,
                      likelihood, impact, owner, treatment, status, created_by,
                      created_at, updated_at
            """,
            (status, treatment, risk_id),
        ).fetchone()
        conn.execute(
            """
            INSERT INTO ai_governance_risk_events
                (risk_id, username, previous_status, new_status, treatment)
            VALUES (%s, %s, %s, %s, %s)
            """,
            (risk_id, username, current["status"], status, treatment),
        )
        return row


def initialize_demand_model_schema() -> None:
    statements = (
        """
        CREATE TABLE IF NOT EXISTS demand_model_runs (
            id TEXT PRIMARY KEY,
            model_version TEXT NOT NULL,
            source_checksum TEXT NOT NULL,
            status TEXT NOT NULL CHECK (status IN ('evaluated', 'failed')),
            created_by TEXT NOT NULL,
            result JSONB NOT NULL,
            model_artifact BYTEA,
            created_at TIMESTAMPTZ NOT NULL DEFAULT now()
        )
        """,
        "CREATE INDEX IF NOT EXISTS demand_model_runs_created_at_idx ON demand_model_runs(created_at DESC)",
        """
        CREATE TABLE IF NOT EXISTS demand_model_predictions (
            id TEXT PRIMARY KEY,
            run_id TEXT NOT NULL REFERENCES demand_model_runs(id) ON DELETE CASCADE,
            period TEXT NOT NULL,
            predicted_value DOUBLE PRECISION NOT NULL,
            unit TEXT NOT NULL,
            traceability JSONB NOT NULL,
            explanation JSONB NOT NULL,
            created_at TIMESTAMPTZ NOT NULL DEFAULT now(),
            UNIQUE (run_id, period)
        )
        """,
        "CREATE INDEX IF NOT EXISTS demand_model_predictions_created_at_idx ON demand_model_predictions(created_at DESC)",
        """
        CREATE TABLE IF NOT EXISTS model_operation_events (
            id BIGSERIAL PRIMARY KEY,
            operation TEXT NOT NULL,
            succeeded BOOLEAN NOT NULL,
            response_time_ms DOUBLE PRECISION NOT NULL CHECK (response_time_ms >= 0),
            username TEXT NOT NULL,
            occurred_at TIMESTAMPTZ NOT NULL DEFAULT now()
        )
        """,
        "CREATE INDEX IF NOT EXISTS model_operation_events_occurred_at_idx ON model_operation_events(occurred_at DESC)",
    )
    with connection() as conn:
        for statement in statements:
            conn.execute(statement)


def migrate_legacy_reports() -> None:
    metadata_path = REPORTS_PATH.parent / "metadata.json"
    if not metadata_path.is_file():
        return
    REPORTS_PATH.mkdir(parents=True, exist_ok=True)
    try:
        metadata = json.loads(metadata_path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as error:
        raise ValueError("No se pudo leer metadata.json del historial de reportes anterior.") from error
    reports = metadata.get("reports") if isinstance(metadata, dict) else None
    if not isinstance(reports, list):
        raise ValueError("metadata.json debe contener una lista 'reports'.")

    with connection() as conn:
        for report in reports:
            if not isinstance(report, dict):
                logger.warning("Se omitió una entrada inválida del historial de reportes anterior.")
                continue
            report_id = report.get("id")
            file_name = report.get("fileName")
            if (
                not isinstance(report_id, str)
                or not isinstance(file_name, str)
                or file_name in {".", ".."}
                or Path(file_name).name != file_name
            ):
                logger.warning("Se omitió un reporte histórico con identificador o nombre de archivo inválido.")
                continue
            source_path = metadata_path.parent / file_name
            target_path = REPORTS_PATH / file_name
            if not source_path.is_file():
                logger.warning("No se encontró el archivo del reporte histórico %s; no se importó.", file_name)
                continue
            created_at_value = report.get("generado")
            try:
                created_at = datetime.fromisoformat(created_at_value.replace("Z", "+00:00"))
            except (AttributeError, ValueError):
                logger.warning("Se omitió el reporte histórico %s por tener una fecha inválida.", report_id)
                continue
            if not target_path.exists():
                shutil.copy2(source_path, target_path)
            conn.execute(
                """
                INSERT INTO generated_reports
                    (id, file_name, name, report_type, period, username, role, sections, size_bytes, created_at)
                VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
                ON CONFLICT (id) DO NOTHING
                """,
                (
                    report_id,
                    file_name,
                    str(report.get("nombre") or report.get("tipo") or "Reporte histórico"),
                    str(report.get("tipo") or "Reporte histórico"),
                    str(report.get("periodo") or "Sin periodo"),
                    str(report.get("usuario") or "Usuario histórico"),
                    "Historial anterior a PostgreSQL",
                    Jsonb([]),
                    target_path.stat().st_size,
                    created_at,
                ),
            )


def _normalized(value: str) -> str:
    decomposed = unicodedata.normalize("NFD", value)
    return "".join(char for char in decomposed if unicodedata.category(char) != "Mn").strip().upper()


def _csv_delimiter(sample: str) -> str:
    try:
        return csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
    except csv.Error:
        first_line = next((line for line in sample.splitlines() if line.strip()), "")
        return max((",", ";", "\t", "|"), key=first_line.count)


def _read_csv_dataset(path: Path) -> list[dict[str, str]]:
    raw_content = path.read_bytes()
    try:
        content = raw_content.decode("utf-8-sig")
    except UnicodeDecodeError:
        # Algunos CSV de fuentes públicas se exportan en Windows-1252 (p. ej., ó = 0xF3).
        content = raw_content.decode("cp1252")
    lines = [line for line in content.splitlines() if line.strip()]
    if not lines:
        return []

    sample = "\n".join(lines[:20])
    delimiter = _csv_delimiter(sample)

    header_index = 0
    header: list[str] = []
    for index, line in enumerate(lines[:20]):
        parsed = next(csv.reader([line], delimiter=delimiter))
        if sum(bool(cell.strip()) for cell in parsed) > sum(bool(cell.strip()) for cell in header):
            header_index = index
            header = parsed

    columns = [value.strip().strip('"') for value in header]
    rows = []
    for line in lines[header_index + 1:]:
        values = next(csv.reader([line], delimiter=delimiter))
        rows.append({column: values[index].strip() if index < len(values) else "" for index, column in enumerate(columns)})
    return rows


def _normalize_complaint_dates(rows: list[dict[str, str]]) -> list[dict[str, str]]:
    normalized_rows = []
    for row in rows:
        normalized = dict(row)
        value = normalized.get("FECHA DE REGISTRO", "").strip()
        try:
            normalized["FECHA DE REGISTRO"] = datetime.strptime(value, "%d/%m/%Y").date().isoformat()
        except ValueError:
            logger.warning("Se conservó una fecha SINADA alternativa con formato inesperado: %s", value)
        normalized_rows.append(normalized)
    return normalized_rows


def _parse_determination_workbook(path: Path, year: int) -> list[dict[str, Any]]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    labels = ("Ene", "Feb", "Mar", "Abr", "May", "Jun", "Jul", "Ago", "Sep", "Oct", "Nov", "Dic")
    monthly = []
    for month_number, label in enumerate(labels, start=1):
        sheet_name = label if label in workbook.sheetnames else "Set" if label == "Sep" and "Set" in workbook.sheetnames else None
        if sheet_name is None:
            continue
        sheet = workbook[sheet_name]
        period_is_2026 = any(
            isinstance(value, str) and "Correspondiente a" in value and str(year) in value
            for row in sheet.iter_rows(values_only=True)
            for value in row
        )
        if not period_is_2026:
            continue
        total_row = next(
            (row for row in sheet.iter_rows(values_only=True) if row and row[0] == "Total General"),
            None,
        )
        if not total_row or len(total_row) <= 15:
            continue
        production_value = _number(total_row[14])
        canon_value = _number(total_row[15])
        if production_value <= 0 or canon_value <= 0:
            continue
        monthly.append({
            "year": year,
            "month": month_number,
            "label": label,
            "productionValueUsd": production_value,
            "canonAndParticipationUsd": canon_value,
        })
    workbook.close()
    return monthly


def _number(value: Any) -> float:
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str):
        try:
            return float(value.replace(",", "").strip())
        except ValueError:
            return 0.0
    return 0.0


def _parse_supervision_workbook(path: Path) -> list[dict[str, Any]]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    monthly: dict[tuple[int, int], dict[str, Any]] = {}
    month_labels = ("Ene", "Feb", "Mar", "Abr", "May", "Jun", "Jul", "Ago", "Sep", "Oct", "Nov", "Dic")
    for sheet_name, value_key, year_header, month_header, value_header in (
        ("MNHL1_Supervisones", "supervisions", "ANO", "MES_NUMERO", "CANTIDAD"),
        ("MNHL1_PAS", "pas", "ANO", "MES", "CANTIDAD_PAS"),
    ):
        if sheet_name not in workbook.sheetnames:
            continue
        rows = workbook[sheet_name].iter_rows(values_only=True)
        header = next(rows, ())
        header_indices = {_normalized(str(value)): index for index, value in enumerate(header) if value is not None}
        year_index = header_indices.get(year_header, -1)
        month_index = header_indices.get(month_header, -1)
        value_index = header_indices.get(value_header, -1)
        date_index = header_indices.get("FECHA", -1)
        if min(year_index, month_index, value_index) < 0:
            continue

        for row in rows:
            try:
                year = int(row[year_index])
            except (TypeError, ValueError):
                continue
            if year < 2018 or year > date.today().year:
                continue
            month_name = _normalized(str(row[month_index]))
            try:
                month_number = int(row[month_index])
            except (TypeError, ValueError):
                month_number = next((index for index, label in enumerate(month_labels, start=1) if _normalized(label) == month_name), 0)
            if not month_number and date_index >= 0:
                month_date = row[date_index]
                if isinstance(month_date, datetime):
                    month_number = month_date.month
                elif isinstance(month_date, date):
                    month_number = month_date.month
            if not month_number:
                continue
            record = monthly.setdefault((year, month_number), {
                "year": year,
                "month": month_number,
                "label": month_labels[month_number - 1],
                "supervisions": 0,
                "pas": 0,
            })
            record[value_key] += _number(row[value_index])
    workbook.close()
    return [monthly[key] for key in sorted(monthly)]


def _dataset_source_needs_refresh(filename: str) -> bool:
    static_sources = set(DATASET_FILES.values()) | {
        ALTERNATE_COMPLAINTS_FILE,
        *REGISTRY_FILES,
        *ENVIRONMENTAL_FILES,
        PRICE_FILE,
        MINORIST_PRICE_FILE,
        "MNHL1-PAS HL.xlsx",
    }
    if filename in static_sources:
        return True
    if filename.startswith("Determin") and filename.lower().endswith(".xlsx") and re.search(r"20\d{2}", filename):
        return True
    return any(filename.startswith(prefix) for prefix in OPERATIONAL_PREFIXES)


def import_source_datasets(
    changed_source_files: set[str] | None = None,
    source_checksums: dict[str, str] | None = None,
) -> None:
    if changed_source_files is not None:
        with connection() as conn:
            stored_keys = {
                row["dataset_key"]
                for row in conn.execute("SELECT dataset_key FROM sector_datasets").fetchall()
            }
        expected_keys = set(DATASET_FILES) | {
            "hydrocarbonData", "registryData", "environmentalData", "priceObservations", "operationalData",
        }
        data_sources_changed = any(_dataset_source_needs_refresh(name) for name in changed_source_files)
        if not data_sources_changed and expected_keys.issubset(stored_keys):
            logger.info("Las fuentes sectoriales no cambiaron; se reutilizan los datasets de PostgreSQL.")
            return

    all_sources: dict[str, dict[str, Any]] = {}
    source_hashes = dict(source_checksums or {})

    def checksum_for(path: Path) -> str:
        checksum = source_hashes.get(path.name)
        if checksum is None:
            checksum = hashlib.sha256(path.read_bytes()).hexdigest()
            source_hashes[path.name] = checksum
        return checksum

    for dataset_key, filename in DATASET_FILES.items():
        path = IMPORTS_PATH / filename
        if not path.is_file():
            raise FileNotFoundError(f"No se encontró el archivo fuente {filename}.")
        all_sources[filename] = {"rows": _read_csv_dataset(path)}
        checksum_for(path)
    alternate_complaints_path = IMPORTS_PATH / ALTERNATE_COMPLAINTS_FILE
    if not alternate_complaints_path.is_file():
        raise FileNotFoundError(f"No se encontró el archivo fuente {ALTERNATE_COMPLAINTS_FILE}.")
    all_sources[ALTERNATE_COMPLAINTS_FILE] = {
        "rows": _normalize_complaint_dates(_read_csv_dataset(alternate_complaints_path))
    }
    checksum_for(alternate_complaints_path)

    determination_paths = sorted(IMPORTS_PATH.glob("Determin*20*.xlsx"))
    if not determination_paths:
        raise FileNotFoundError("No se encontraron libros anuales de determinación.")
    workbook_sources = {
        **{path.name: path for path in determination_paths},
        "MNHL1-PAS HL.xlsx": IMPORTS_PATH / "MNHL1-PAS HL.xlsx",
    }
    for source_key, path in workbook_sources.items():
        if not path.is_file():
            raise FileNotFoundError(f"No se encontró el archivo fuente {path.name}.")
        checksum_for(path)

    determinations = []
    for path in determination_paths:
        year_match = re.search(r"20\d{2}", path.name)
        if year_match:
            determinations.extend(_parse_determination_workbook(path, int(year_match.group())))
    hydrocarbon_data = {
        "determinations": sorted(determinations, key=lambda row: (row["year"], row["month"])),
        "supervision": _parse_supervision_workbook(workbook_sources["MNHL1-PAS HL.xlsx"]),
    }
    if not hydrocarbon_data["determinations"] or not hydrocarbon_data["supervision"]:
        raise ValueError("Las fuentes anuales de determinación y supervisión no contienen datos válidos.")

    all_sources["Determinación+2026.xlsx"] = {"rows": hydrocarbon_data["determinations"]}
    all_sources["MNHL1-PAS HL.xlsx"] = {"rows": hydrocarbon_data["supervision"]}
    all_sources["hydrocarbonData"] = {"rows": hydrocarbon_data}
    for dataset_key, filename in DATASET_FILES.items():
        all_sources[dataset_key] = {"rows": all_sources[filename]["rows"]}
        source_hashes[dataset_key] = source_hashes[filename]
    primary_complaints = all_sources[DATASET_FILES["complaints"]]["rows"]
    alternate_complaints = all_sources[ALTERNATE_COMPLAINTS_FILE]["rows"]
    existing_complaint_ids = {row.get("CODIGO SINADA", "").strip() for row in primary_complaints}
    all_sources["complaints"] = {
        "rows": primary_complaints + [
            row for row in alternate_complaints
            if row.get("CODIGO SINADA", "").strip() not in existing_complaint_ids
        ]
    }
    source_hashes["complaints"] = hashlib.sha256(
        (source_hashes[DATASET_FILES["complaints"]] + source_hashes[ALTERNATE_COMPLAINTS_FILE]).encode()
    ).hexdigest()
    source_hashes["hydrocarbonData"] = hashlib.sha256(
        "".join(source_hashes[name] for name in [*workbook_sources]).encode()
    ).hexdigest()
    operational_paths = [
        path for prefix in OPERATIONAL_PREFIXES
        for path in IMPORTS_PATH.glob(f"{prefix}*.xlsx")
    ]
    if len(operational_paths) != len(OPERATIONAL_PREFIXES):
        raise FileNotFoundError("Faltan libros Excel operativos requeridos para los indicadores sectoriales.")
    for path in operational_paths:
        checksum_for(path)
    source_hashes["operationalData"] = hashlib.sha256(
        "".join(f"{path.name}:{source_hashes[path.name]}" for path in sorted(operational_paths)).encode()
    ).hexdigest()
    imported_data = build_imported_data(IMPORTS_PATH)
    imported_source_files = {
        "registryData": [source["fileName"] for source in imported_data["registryData"]["sources"]],
        "environmentalData": [source["fileName"] for source in imported_data["environmentalData"]["sources"]],
        "priceObservations": (
            [imported_data["priceObservations"]["sourceFile"]]
            + (
                [imported_data["priceObservations"]["minoristRegistry"]["sourceFile"]]
                if (IMPORTS_PATH / imported_data["priceObservations"]["minoristRegistry"]["sourceFile"]).is_file()
                else []
            )
            if imported_data["priceObservations"].get("available")
            else []
        ),
    }
    for dataset_key, source_files in imported_source_files.items():
        file_hashes = []
        for source_file in source_files:
            path = IMPORTS_PATH / source_file
            checksum = checksum_for(path)
            file_hashes.append(f"{source_file}:{checksum}")
        source_hashes[dataset_key] = hashlib.sha256("|".join(file_hashes).encode()).hexdigest()
        all_sources[dataset_key] = {"rows": imported_data[dataset_key]}

    with connection() as conn:
        conn.execute(
            "DELETE FROM sector_datasets WHERE dataset_key IN (%s, %s)",
            ("Determination+2026.xlsx", "Supervision+PAS.xlsx"),
        )
        for dataset_key, dataset in all_sources.items():
            source_files = (
                [DATASET_FILES[dataset_key]]
                if dataset_key in DATASET_FILES
                else imported_source_files[dataset_key]
                if dataset_key in imported_source_files
                else [dataset_key] if dataset_key in source_hashes
                else [filename for filename in dataset if filename.endswith((".csv", ".xlsx"))]
            )
            if dataset_key == "complaints":
                source_files = [DATASET_FILES["complaints"], ALTERNATE_COMPLAINTS_FILE]
            if dataset_key == "hydrocarbonData":
                source_files = [path.name for path in determination_paths] + ["MNHL1-PAS HL.xlsx"]
            checksum = source_hashes.get(dataset_key)
            if checksum is None:
                source_file = source_files[0]
                checksum = source_hashes[source_file]
            current = conn.execute(
                "SELECT checksum FROM sector_datasets WHERE dataset_key = %s",
                (dataset_key,),
            ).fetchone()
            if current and current["checksum"] == checksum:
                continue
            conn.execute(
                """
                INSERT INTO sector_datasets (dataset_key, source_files, checksum, rows, updated_at)
                VALUES (%s, %s, %s, %s, now())
                ON CONFLICT (dataset_key) DO UPDATE SET
                    source_files = EXCLUDED.source_files,
                    checksum = EXCLUDED.checksum,
                    rows = EXCLUDED.rows,
                    updated_at = now()
                """,
                (dataset_key, Jsonb(source_files), checksum, Jsonb(dataset["rows"])),
            )
        current = conn.execute(
            "SELECT checksum FROM sector_datasets WHERE dataset_key = %s",
            ("operationalData",),
        ).fetchone()
        if not current or current["checksum"] != source_hashes["operationalData"]:
            operational_data = build_operational_data(IMPORTS_PATH)
            conn.execute(
                """
                INSERT INTO sector_datasets (dataset_key, source_files, checksum, rows, updated_at)
                VALUES (%s, %s, %s, %s, now())
                ON CONFLICT (dataset_key) DO UPDATE SET
                    source_files = EXCLUDED.source_files,
                    checksum = EXCLUDED.checksum,
                    rows = EXCLUDED.rows,
                    updated_at = now()
                """,
                (
                    "operationalData",
                    Jsonb([path.name for path in sorted(operational_paths)]),
                    source_hashes["operationalData"],
                    Jsonb(operational_data),
                ),
            )


def import_source_files() -> tuple[set[str], dict[str, str]]:
    if not IMPORTS_PATH.is_dir():
        raise FileNotFoundError(f"No se encontró el directorio de fuentes {IMPORTS_PATH}.")
    paths = sorted(path for path in IMPORTS_PATH.iterdir() if path.is_file())
    if not paths:
        raise FileNotFoundError("El directorio de fuentes no contiene archivos para importar.")
    imported_names = [path.name for path in paths]

    changed_names: set[str] = set()
    checksums: dict[str, str] = {}
    with connection() as conn:
        current_rows = conn.execute(
            "SELECT file_name, checksum, size_bytes, modified_time_ns FROM source_files"
        ).fetchall()
        current_by_name = {row["file_name"]: row for row in current_rows}
        for path in paths:
            stat = path.stat()
            current = current_by_name.get(path.name)
            if (
                current
                and current["size_bytes"] == stat.st_size
                and current["modified_time_ns"] == stat.st_mtime_ns
            ):
                checksums[path.name] = current["checksum"]
                continue

            # Los originales pueden superar ampliamente el tamaño razonable para
            # transferirlos como BYTEA en una conexión administrada. El contenedor
            # ya conserva estos archivos en modo lectura; PostgreSQL mantiene su
            # catálogo, huella, tamaño y los datos estructurados derivados.
            file_hash = hashlib.sha256()
            with path.open("rb") as source_file:
                for chunk in iter(lambda: source_file.read(1024 * 1024), b""):
                    file_hash.update(chunk)
            checksum = file_hash.hexdigest()
            checksums[path.name] = checksum
            if current and current["checksum"] == checksum:
                conn.execute(
                    "UPDATE source_files SET size_bytes = %s, modified_time_ns = %s WHERE file_name = %s",
                    (stat.st_size, stat.st_mtime_ns, path.name),
                )
                continue
            changed_names.add(path.name)
            file_type = path.suffix.lower().lstrip(".")
            conn.execute(
                """
                INSERT INTO source_files
                    (file_name, file_type, checksum, size_bytes, modified_time_ns, content, updated_at)
                VALUES (%s, %s, %s, %s, %s, NULL, now())
                ON CONFLICT (file_name) DO UPDATE SET
                    file_type = EXCLUDED.file_type,
                    checksum = EXCLUDED.checksum,
                    size_bytes = EXCLUDED.size_bytes,
                    modified_time_ns = EXCLUDED.modified_time_ns,
                    content = NULL,
                    updated_at = now()
                """,
                (path.name, file_type, checksum, stat.st_size, stat.st_mtime_ns),
            )
        stale_names = set(current_by_name) - set(imported_names)
        changed_names.update(stale_names)
        conn.execute("DELETE FROM source_files WHERE NOT (file_name = ANY(%s))", (imported_names,))
    return changed_names, checksums


def list_source_files() -> list[dict[str, Any]]:
    with connection() as conn:
        records = conn.execute(
            """
            SELECT file_name, file_type, size_bytes, updated_at
            FROM source_files ORDER BY file_type, file_name
            """
        ).fetchall()
    return [
        {
            "fileName": row["file_name"],
            "fileType": row["file_type"],
            "sizeBytes": row["size_bytes"],
            "updatedAt": row["updated_at"].isoformat(),
        }
        for row in records
    ]


def get_source_file(file_name: str) -> dict[str, Any] | None:
    if not file_name or Path(file_name).name != file_name or file_name in {".", ".."}:
        return None
    with connection() as conn:
        record = conn.execute(
            "SELECT file_name, file_type, size_bytes, content FROM source_files WHERE file_name = %s",
            (file_name,),
        ).fetchone()
    if not record:
        return None
    content = bytes(record["content"]) if record["content"] is not None else None
    if content is None:
        source_path = IMPORTS_PATH / file_name
        if not source_path.is_file():
            return None
        content = source_path.read_bytes()
    media_type = mimetypes.guess_type(record["file_name"])[0] or "application/octet-stream"
    return {
        "fileName": record["file_name"],
        "fileType": record["file_type"],
        "sizeBytes": record["size_bytes"],
        "content": content,
        "mediaType": media_type,
    }


def preview_source_file(file_name: str, sheet_name: str | None = None) -> dict[str, Any] | None:
    source = get_source_file(file_name)
    if source is None:
        return None
    if source["fileType"] == "pdf":
        return {
            "fileName": file_name,
            "fileType": "pdf",
            "previewAvailable": False,
            "sheets": [],
            "rows": [],
            "totalRows": 0,
        }

    if source["fileType"] == "xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(source["content"]), read_only=True, data_only=True)
        try:
            sheet_names = workbook.sheetnames
            selected_sheet = sheet_name or (sheet_names[0] if sheet_names else None)
            if selected_sheet not in sheet_names:
                raise ValueError(f"La hoja solicitada no existe en {file_name}.")
            worksheet = workbook[selected_sheet]
            rows = [
                [_preview_cell(value) for value in row]
                for row in worksheet.iter_rows(max_row=100, values_only=True)
            ]
            while rows and not any(value is not None for value in rows[-1]):
                rows.pop()
            return {
                "fileName": file_name,
                "fileType": "xlsx",
                "previewAvailable": True,
                "sheets": sheet_names,
                "selectedSheet": selected_sheet,
                "rows": rows,
                "totalRows": worksheet.max_row,
            }
        finally:
            workbook.close()

    if source["fileType"] == "csv":
        content = source["content"].decode("utf-8-sig")
        lines = [line for line in content.splitlines() if line.strip()]
        sample = "\n".join(lines[:20])
        delimiter = _csv_delimiter(sample)
        parsed_rows = list(csv.reader(io.StringIO(content), delimiter=delimiter))
        header_index = max(
            range(min(20, len(parsed_rows))),
            key=lambda index: sum(bool(value.strip()) for value in parsed_rows[index]),
            default=0,
        )
        rows = [
            [_preview_cell(value) for value in row]
            for row in parsed_rows[header_index + 1:]
            if any(value.strip() for value in row)
        ]
        return {
            "fileName": file_name,
            "fileType": "csv",
            "previewAvailable": True,
            "sheets": [],
            "columns": [_preview_cell(value) for value in parsed_rows[header_index]],
            "rows": rows[:100],
            "totalRows": len(rows),
        }

    return {
        "fileName": file_name,
        "fileType": source["fileType"],
        "previewAvailable": False,
        "sheets": [],
        "rows": [],
        "totalRows": 0,
    }


def _preview_cell(value: Any) -> str | int | float | bool | None:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return str(value)


def hash_password(password: str) -> str:
    salt = secrets.token_bytes(16)
    digest = hashlib.pbkdf2_hmac("sha256", password.encode(), salt, 310_000)
    return f"pbkdf2_sha256$310000${salt.hex()}${digest.hex()}"


def verify_password(password: str, encoded: str) -> bool:
    try:
        algorithm, iterations, salt_hex, digest_hex = encoded.split("$")
        if algorithm != "pbkdf2_sha256":
            return False
        expected = bytes.fromhex(digest_hex)
        actual = hashlib.pbkdf2_hmac("sha256", password.encode(), bytes.fromhex(salt_hex), int(iterations))
        return secrets.compare_digest(actual, expected)
    except (ValueError, TypeError):
        return False


def seed_initial_supervisor() -> None:
    username = os.environ.get("SEED_ADMIN_USERNAME", "").strip().lower()
    password = os.environ.get("SEED_ADMIN_PASSWORD", "")
    display_name = os.environ.get("SEED_ADMIN_NAME", "Supervisor de Cumplimiento").strip()
    if not username and not password:
        with connection() as conn:
            existing = conn.execute("SELECT 1 FROM managed_users LIMIT 1").fetchone()
        if existing:
            return
        raise ValueError("Configura SEED_ADMIN_USERNAME y SEED_ADMIN_PASSWORD para crear el supervisor inicial.")
    if not username or len(username) < 3 or len(password) < 12:
        raise ValueError("SEED_ADMIN_USERNAME y una contraseña SEED_ADMIN_PASSWORD de al menos 12 caracteres son obligatorios.")
    if not display_name:
        raise ValueError("SEED_ADMIN_NAME no puede estar vacío.")

    with connection() as conn:
        existing = conn.execute("SELECT 1 FROM managed_users WHERE username = %s", (username,)).fetchone()
        if existing:
            return
        conn.execute(
            """
            INSERT INTO managed_users (username, display_name, password_hash, role, permissions)
            VALUES (%s, %s, %s, %s, %s)
            """,
            (
                username,
                display_name,
                hash_password(password),
                "Supervisor de Cumplimiento",
                Jsonb(PERMISSIONS_BY_ROLE["Supervisor de Cumplimiento"]),
            ),
        )


def public_user(row: dict[str, Any]) -> dict[str, Any]:
    return {
        "username": row["username"],
        "name": row["display_name"],
        "role": row["role"],
        "permissions": row["permissions"],
        "active": row["active"],
    }


def token_digest(token: str) -> str:
    return hashlib.sha256(token.encode()).hexdigest()


def create_session(username: str) -> str:
    token = secrets.token_urlsafe(32)
    with connection() as conn:
        conn.execute("DELETE FROM auth_sessions WHERE expires_at <= now()")
        conn.execute(
            "INSERT INTO auth_sessions (token_hash, username, expires_at) VALUES (%s, %s, %s)",
            (token_digest(token), username, datetime.now(timezone.utc) + SESSION_TTL),
        )
    return token


def lookup_session(token: str) -> dict[str, Any] | None:
    with connection() as conn:
        return conn.execute(
            """
            SELECT u.username, u.display_name, u.role, u.permissions, u.active
            FROM auth_sessions s
            JOIN managed_users u ON u.username = s.username
            WHERE s.token_hash = %s AND s.expires_at > now() AND u.active
            """,
            (token_digest(token),),
        ).fetchone()


def load_datasets() -> dict[str, Any]:
    with connection() as conn:
        records = conn.execute("SELECT dataset_key, rows, source_files, updated_at FROM sector_datasets ORDER BY dataset_key").fetchall()
    return {
        row["dataset_key"]: {
            "rows": row["rows"],
            "sourceFiles": row["source_files"],
            "updatedAt": row["updated_at"].isoformat(),
        }
        for row in records
    }


def load_imported_data() -> dict[str, Any]:
    dataset_keys = ("registryData", "environmentalData", "priceObservations")
    with connection() as conn:
        records = conn.execute(
            "SELECT dataset_key, rows FROM sector_datasets WHERE dataset_key = ANY(%s)",
            (list(dataset_keys),),
        ).fetchall()
    return {row["dataset_key"]: row["rows"] for row in records}


def save_demand_model_run(result: dict[str, Any], model_artifact: bytes, username: str) -> None:
    with connection() as conn:
        conn.execute(
            """
            INSERT INTO demand_model_runs
                (id, model_version, source_checksum, status, created_by, result, model_artifact)
            VALUES (%s, %s, %s, 'evaluated', %s, %s, %s)
            """,
            (
                result["runId"],
                result["modelVersion"],
                result["sourceChecksum"],
                username,
                Jsonb(result),
                model_artifact,
            ),
        )


def load_demand_model_run(run_id: str | None = None) -> dict[str, Any] | None:
    with connection() as conn:
        if run_id is None:
            row = conn.execute(
                """
                SELECT id, source_checksum, result, model_artifact
                FROM demand_model_runs
                WHERE status = 'evaluated' AND model_artifact IS NOT NULL
                ORDER BY created_at DESC LIMIT 1
                """
            ).fetchone()
        else:
            row = conn.execute(
                """
                SELECT id, source_checksum, result, model_artifact
                FROM demand_model_runs
                WHERE id = %s AND status = 'evaluated' AND model_artifact IS NOT NULL
                """,
                (run_id,),
            ).fetchone()
    return row


def list_demand_model_runs(limit: int = 20) -> list[dict[str, Any]]:
    with connection() as conn:
        return conn.execute(
            """
            SELECT result
            FROM demand_model_runs
            ORDER BY created_at DESC
            LIMIT %s
            """,
            (limit,),
        ).fetchall()


def save_demand_prediction(
    *,
    prediction_id: str,
    run_id: str,
    period: str,
    value: float,
    unit: str,
    traceability: dict[str, Any],
    explanation: dict[str, Any],
) -> None:
    with connection() as conn:
        conn.execute(
            """
            INSERT INTO demand_model_predictions
                (id, run_id, period, predicted_value, unit, traceability, explanation)
            VALUES (%s, %s, %s, %s, %s, %s, %s)
            ON CONFLICT (run_id, period) DO UPDATE SET
                predicted_value = EXCLUDED.predicted_value,
                unit = EXCLUDED.unit,
                traceability = EXCLUDED.traceability,
                explanation = EXCLUDED.explanation,
                created_at = now()
            """,
            (
                prediction_id,
                run_id,
                period,
                value,
                unit,
                Jsonb(traceability),
                Jsonb(explanation),
            ),
        )


def list_demand_model_predictions(limit: int = 100) -> list[dict[str, Any]]:
    with connection() as conn:
        return conn.execute(
            """
            SELECT p.id, p.run_id, r.model_version, p.period, p.predicted_value, p.unit,
                   p.traceability, p.explanation, p.created_at
            FROM demand_model_predictions p
            JOIN demand_model_runs r ON r.id = p.run_id
            ORDER BY p.created_at DESC
            LIMIT %s
            """,
            (limit,),
        ).fetchall()


def record_model_operation(operation: str, succeeded: bool, response_time_ms: float, username: str) -> None:
    with connection() as conn:
        conn.execute(
            """
            INSERT INTO model_operation_events (operation, succeeded, response_time_ms, username)
            VALUES (%s, %s, %s, %s)
            """,
            (operation, succeeded, response_time_ms, username),
        )


def get_model_operation_metrics() -> dict[str, Any]:
    with connection() as conn:
        result = conn.execute(
            """
            SELECT COUNT(*) AS executed,
                   COUNT(*) FILTER (WHERE succeeded) AS succeeded,
                   AVG(response_time_ms) AS mean_response_time_ms
            FROM model_operation_events
            """
        ).fetchone()
        prediction_metrics = conn.execute(
            """
            SELECT COUNT(*) AS total,
                   COUNT(*) FILTER (
                       WHERE NULLIF(traceability ->> 'runId', '') IS NOT NULL
                         AND NULLIF(traceability ->> 'modelVersion', '') IS NOT NULL
                         AND NULLIF(traceability ->> 'targetPeriod', '') IS NOT NULL
                         AND NULLIF(traceability ->> 'source', '') IS NOT NULL
                         AND NULLIF(traceability ->> 'sourcePeriod', '') IS NOT NULL
                         AND NULLIF(traceability ->> 'generatedAt', '') IS NOT NULL
                         AND NULLIF(traceability ->> 'requestedBy', '') IS NOT NULL
                         AND NULLIF(traceability ->> 'unit', '') IS NOT NULL
                   ) AS traceable,
                   COUNT(*) FILTER (WHERE explanation ->> 'status' = 'available') AS explained
            FROM demand_model_predictions
            """
        ).fetchone()
        control_metrics = conn.execute(
            """
            SELECT COUNT(*) AS total,
                   COUNT(*) FILTER (WHERE status = 'conforme' AND length(trim(evidence)) > 0) AS conforming
            FROM ai_governance_controls
            WHERE model_scope IN ('demanda', 'ambos')
            """
        ).fetchone()

    total = prediction_metrics["total"]
    traceable = prediction_metrics["traceable"]
    explained = prediction_metrics["explained"]
    return {
        "predictionTraceability": {
            "status": "measured" if total else "pending",
            "totalPredictions": total,
            "completePredictions": traceable,
            "percent": round(traceable / total * 100, 2) if total else None,
        },
        "explainabilityCoverage": {
            "status": "measured" if total else "pending",
            "totalPredictions": total,
            "explainedPredictions": explained,
            "percent": round(explained / total * 100, 2) if total else None,
        },
        "governanceControlsCompliance": {
            "status": "measured" if control_metrics["total"] else "pending",
            "totalControls": control_metrics["total"],
            "controlsWithEvidenceMarkedConforming": control_metrics["conforming"],
            "reason": "Porcentaje de controles que un supervisor marcó conformes con evidencia registrada; autoevaluación del proyecto, no auditoría ni certificación.",
            "percent": round(control_metrics["conforming"] / control_metrics["total"] * 100, 2) if control_metrics["total"] else None,
        },
        "riskDetectionTime": {
            "status": "unavailable",
            "reason": "No hay eventos de riesgo con tiempos verificables de registro y detección.",
            "meanMinutes": None,
        },
        "systemOperations": {
            "status": "measured" if result["executed"] else "pending",
            "executed": result["executed"],
            "succeeded": result["succeeded"],
            "successRatePercent": (
                round(result["succeeded"] / result["executed"] * 100, 2)
                if result["executed"]
                else None
            ),
            "meanResponseTimeMs": (
                round(result["mean_response_time_ms"], 2)
                if result["mean_response_time_ms"] is not None
                else None
            ),
        },
        "securityTestPassRate": {
            "status": "unavailable",
            "reason": "No hay resultados de pruebas de seguridad persistidos en el sistema.",
            "percent": None,
        },
        "aiRiskDetectionRate": {
            "status": "unavailable",
            "reason": "No hay un conjunto de riesgos conocidos etiquetados para medir detección.",
            "percent": None,
        },
        "perceivedUtility": {
            "status": "unavailable",
            "reason": "No hay encuestas reales registradas.",
            "percent": None,
        },
    }


def save_risk_assessment(assessment: dict[str, Any]) -> dict[str, Any]:
    with connection() as conn:
        row = conn.execute(
            """
            INSERT INTO risk_assessments
                (id, scenario_name, inputs, score, risk_level, factors, assessment_type, created_by)
            VALUES (%s, %s, %s, %s, %s, %s, 'simulación', %s)
            RETURNING id, scenario_name, inputs, score, risk_level, factors,
                      assessment_type, created_by, created_at
            """,
            (
                assessment["id"], assessment["scenarioName"], Jsonb(assessment["inputs"]),
                assessment["score"], assessment["riskLevel"], Jsonb(assessment["factors"]),
                assessment["createdBy"],
            ),
        ).fetchone()
        conn.execute(
            """
            INSERT INTO access_events (username, role, action, resource, result)
            VALUES (%s, %s, 'Registro de simulación de riesgo', %s, 'Permitido')
            """,
            (assessment["createdBy"], assessment["createdByRole"], assessment["id"]),
        )
    return _public_risk_assessment(row)


def list_risk_assessments(limit: int = 100) -> list[dict[str, Any]]:
    with connection() as conn:
        rows = conn.execute(
            """
            SELECT id, scenario_name, inputs, score, risk_level, factors,
                   assessment_type, created_by, created_at
            FROM risk_assessments ORDER BY created_at DESC, id DESC LIMIT %s
            """,
            (limit,),
        ).fetchall()
    return [_public_risk_assessment(row) for row in rows]


def _public_risk_assessment(row: dict[str, Any]) -> dict[str, Any]:
    return {
        "id": row["id"],
        "scenarioName": row["scenario_name"],
        "inputs": row["inputs"],
        "score": row["score"],
        "riskLevel": row["risk_level"],
        "factors": row["factors"],
        "assessmentType": row["assessment_type"],
        "createdBy": row["created_by"],
        "createdAt": row["created_at"].isoformat(),
    }

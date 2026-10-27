import re
import unicodedata
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


MONTHS = {
    "ENERO": 1, "FEBRERO": 2, "MARZO": 3, "ABRIL": 4, "MAYO": 5, "JUNIO": 6,
    "JULIO": 7, "AGOSTO": 8, "SEPTIEMBRE": 9, "SETIEMBRE": 9, "OCTUBRE": 10,
    "NOVIEMBRE": 11, "DICIEMBRE": 12,
}


def _normalize(value: Any) -> str:
    decomposed = unicodedata.normalize("NFD", str(value or ""))
    return "".join(char for char in decomposed if unicodedata.category(char) != "Mn").strip().upper()


def _field(row: dict[str, Any], *names: str) -> Any:
    normalized = {
        re.sub(r"[^A-Z0-9]", "", _normalize(key)): value
        for key, value in row.items()
    }
    for name in names:
        key = re.sub(r"[^A-Z0-9]", "", _normalize(name))
        if key in normalized:
            return normalized[key]
    return None


def _number(value: Any) -> float:
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str):
        try:
            return float(value.replace(",", "").strip())
        except ValueError:
            return 0.0
    return 0.0


def _period(value: Any, year: Any = None, month: Any = None) -> str | None:
    if isinstance(value, (date, datetime)):
        return f"{value.year:04d}-{value.month:02d}"
    if isinstance(value, str):
        iso = re.search(r"(20\d{2})[-/](\d{1,2})", value)
        if iso:
            month_number = int(iso.group(2))
            if 1 <= month_number <= 12:
                return f"{int(iso.group(1)):04d}-{month_number:02d}"
        dmy = re.search(r"(\d{1,2})/(\d{1,2})/(20\d{2})", value)
        if dmy:
            month_number = int(dmy.group(2))
            if 1 <= month_number <= 12:
                return f"{int(dmy.group(3)):04d}-{month_number:02d}"
    try:
        year_number = int(year)
    except (TypeError, ValueError):
        return None
    try:
        month_number = int(month)
    except (TypeError, ValueError):
        month_number = MONTHS.get(_normalize(month), 0)
    return f"{year_number:04d}-{month_number:02d}" if 1 <= month_number <= 12 else None


def _workbook_path(root: Path, prefix: str) -> Path:
    matches = sorted(root.glob(f"{prefix}*.xlsx"))
    if len(matches) != 1:
        raise FileNotFoundError(f"Se esperaba un archivo Excel único para {prefix}; encontrados: {len(matches)}.")
    return matches[0]


def _records(path: Path, sheet_name: str, header_row: int = 1, max_rows: int | None = None):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"No existe la hoja {sheet_name} en {path.name}.")
        worksheet = workbook[sheet_name]
        header = next(worksheet.iter_rows(min_row=header_row, max_row=header_row, values_only=True), ())
        names = [str(value).strip() if value is not None else f"column_{index + 1}" for index, value in enumerate(header)]
        end_row = min(worksheet.max_row, max_rows) if max_rows is not None else worksheet.max_row
        for values in worksheet.iter_rows(min_row=header_row + 1, max_row=end_row, values_only=True):
            if any(value is not None for value in values):
                yield {names[index]: value for index, value in enumerate(values) if index < len(names)}
    finally:
        workbook.close()


def _group_sum(rows: list[dict[str, Any]], dimensions: tuple[str, ...], metric: str) -> list[dict[str, Any]]:
    totals: dict[tuple[Any, ...], float] = defaultdict(float)
    for row in rows:
        key = tuple(row[dimension] for dimension in dimensions)
        totals[key] += row[metric]
    return [
        {**dict(zip(dimensions, key)), metric: value}
        for key, value in sorted(totals.items(), key=lambda item: item[0])
    ]


def _read_gas(root: Path) -> dict[str, Any]:
    production_path = _workbook_path(root, "MNGN2-3-")
    production = []
    for row in _records(production_path, "GN humedo-seco"):
        period = _period(_field(row, "fecha"))
        if not period:
            continue
        facilities = [key for key in row if _normalize(key) not in {"FECHA", "TIPO"}]
        production.append({
            "period": period,
            "type": str(_field(row, "Tipo") or "").strip(),
            "volume": sum(_number(row[key]) for key in facilities),
        })

    exports = []
    export_path = _workbook_path(root, "MNGN4-")
    for row in _records(export_path, "MNGN4"):
        period = _period(_field(row, "Fecha de Embarque"))
        if period:
            exports.append({"period": period, "volumeCubicMeters": _number(_field(row, "Metros Cúbicos"))})

    transport = []
    transport_path = _workbook_path(root, "MNGN5-")
    for row in _records(transport_path, "BD_DUCTO DE GAS"):
        period = _period(_field(row, "FECHA"))
        if period:
            transport.append({
                "period": period,
                "packedVolume": _number(_field(row, "Volumen Empacado TGP")),
                "injectedVolume": _number(_field(row, "Volumen Inyectado Malvinas")),
            })
    transport_monthly = _group_sum(transport, ("period",), "packedVolume")
    injection_monthly = {row["period"]: row["injectedVolume"] for row in _group_sum(transport, ("period",), "injectedVolume")}
    for row in transport_monthly:
        row["injectedVolume"] = injection_monthly.get(row["period"], 0)

    distribution = []
    distribution_path = _workbook_path(root, "MNGN6-")
    for row in _records(distribution_path, "dist_cli_vol"):
        period = _period(None, _field(row, "Año"), _field(row, "Mes"))
        if period:
            distribution.append({
                "period": period,
                "concession": str(_field(row, "Concesión") or "").strip(),
                "volume": _number(_field(row, "Volumen de Gas Natural distribuido")),
                "customers": _number(_field(row, "Número de clientes")),
            })
    distribution_volume = _group_sum(distribution, ("period", "concession"), "volume")
    distribution_customers = {
        (row["period"], row["concession"]): row["customers"]
        for row in _group_sum(distribution, ("period", "concession"), "customers")
    }
    for row in distribution_volume:
        row["customers"] = distribution_customers.get((row["period"], row["concession"]), 0)

    network = []
    network_path = _workbook_path(root, "MNGN7-")
    for row in _records(network_path, "T01 - Expansión de redes"):
        period = _period(_field(row, "Mes"))
        if period:
            network.append({
                "period": period,
                "concession": str(_field(row, "Concesión") or "").strip(),
                "zone": str(_field(row, "Zona") or "").strip(),
                "kilometers": _number(_field(row, "Total kilómetros")),
            })
    network_monthly = _group_sum(network, ("period", "concession"), "kilometers")

    tariffs = []
    tariff_path = _workbook_path(root, "MNGN8-")
    workbook = load_workbook(tariff_path, read_only=True, data_only=True)
    try:
        for sheet_name in workbook.sheetnames:
            if _normalize(sheet_name) == "RESOLUCIONES":
                continue
            worksheet = workbook[sheet_name]
            header_row = 5 if _normalize(sheet_name) == "CALIDDA" else 3
            header_rows = [
                next(worksheet.iter_rows(min_row=index, max_row=index, values_only=True), ())
                for index in (2, 3, 4)
            ]
            rows = []
            for values in worksheet.iter_rows(min_row=header_row + 1, max_row=min(worksheet.max_row, 1500), values_only=True):
                if not values or not values[0]:
                    continue
                rates = []
                for column_index, value in enumerate(values[5:], start=5):
                    if value is None:
                        continue
                    try:
                        rate = float(str(value).replace(",", "").strip())
                    except ValueError:
                        continue
                    label = next(
                        (
                            str(header_rows[row_index][column_index]).strip()
                            for row_index in (2, 1, 0)
                            if column_index < len(header_rows[row_index])
                            and header_rows[row_index][column_index] is not None
                        ),
                        f"Componente {column_index + 1}",
                    )
                    rates.append({"label": label, "value": rate})
                rows.append({
                    "period": _period(values[0]),
                    "category": str(values[3] or "").strip() if len(values) > 3 else "",
                    "range": str((values[4] if len(values) > 4 else "") or "").strip(),
                    "rates": rates,
                })
            dated_rows = [
                (row["period"], row)
                for row in rows
                if row["period"]
            ]
            if not dated_rows:
                continue
            latest_period = max(period for period, _ in dated_rows)
            latest_rows = [
                row for period, row in dated_rows if period == latest_period
            ]
            tariffs.append({
                "concession": sheet_name,
                "period": latest_period,
                "categories": [
                    {
                        "category": row["category"],
                        "range": row["range"],
                        "rates": row["rates"],
                    }
                    for row in latest_rows
                    if row["category"]
                ],
            })
    finally:
        workbook.close()

    return {
        "production": production,
        "exports": _group_sum(exports, ("period",), "volumeCubicMeters"),
        "transport": transport_monthly,
        "distribution": distribution_volume,
        "network": network_monthly,
        "tariffs": tariffs,
    }


def _read_fuels(root: Path) -> dict[str, Any]:
    fines = []
    fines_path = _workbook_path(root, "MNHL2-")
    for row in _records(fines_path, "MNHL2-Multas unidades operativa"):
        period = _period(None, _field(row, "Año"), _field(row, "Mes"))
        if period:
            fines.append({
                "period": period,
                "count": _number(_field(row, "Cantidad de multas")),
                "uit": _number(_field(row, "Multa en UIT")),
            })
    fine_counts = _group_sum(fines, ("period",), "count")
    fine_amounts = {row["period"]: row["uit"] for row in _group_sum(fines, ("period",), "uit")}
    for row in fine_counts:
        row["uit"] = fine_amounts.get(row["period"], 0)

    prices = []
    prices_path = _workbook_path(root, "MNHL6-")
    for row in _records(prices_path, "MNHL6-Precios Combustibles"):
        period = _period(_field(row, "SEMANA"))
        price = _number(_field(row, "PRECIO PROMEDIO"))
        if period and price > 0:
            prices.append({"period": period, "product": str(_field(row, "COMBUSTIBLE") or "").strip(), "price": price})
    price_totals: dict[tuple[str, str], list[float]] = defaultdict(list)
    for row in prices:
        price_totals[(row["period"], row["product"])].append(row["price"])
    monthly_prices = [
        {"period": period, "product": product, "averagePrice": sum(values) / len(values)}
        for (period, product), values in sorted(price_totals.items())
    ]

    sales = []
    sales_path = _workbook_path(root, "MNHL7-")
    for row in _records(sales_path, "MNHL7-Ventas Comb"):
        period = _period(_field(row, "FCHA_DESPACHO"))
        if period:
            sales.append({
                "period": period,
                "product": str(_field(row, "EQ_PRODUCTO") or "").strip(),
                "gallons": _number(_field(row, "GALONES")),
            })

    glp_prices = []
    glp_price_path = _workbook_path(root, "MNHL8-")
    for row in _records(glp_price_path, "MNHL8-Precios GLP"):
        period = _period(_field(row, "SEMANA"))
        price = _number(_field(row, "PRECIO PROMEDIO"))
        if period and price > 0:
            glp_prices.append({"period": period, "price": price})
    glp_by_period: dict[str, list[float]] = defaultdict(list)
    for row in glp_prices:
        glp_by_period[row["period"]].append(row["price"])
    glp_average_prices = [
        {"period": period, "averagePrice": sum(values) / len(values)}
        for period, values in sorted(glp_by_period.items())
    ]

    glp_sales = []
    glp_sales_path = _workbook_path(root, "MNHL9-")
    for row in _records(glp_sales_path, "VENTAS DE GLP REGISTRADAS"):
        period = _period(None, _field(row, "AÑO"), _field(row, "MES"))
        if period:
            glp_sales.append({
                "period": period,
                "product": str(_field(row, "PRODUCTO") or "").strip(),
                "barrels": _number(_field(row, "BARRILES")),
            })

    trade = []
    trade_path = _workbook_path(root, "MNHL4-")
    for sheet_name, direction in (("BD_Exportacion", "Exportación"), ("BD_Importacion", "Importación")):
        for row in _records(trade_path, sheet_name):
            period = _period(_field(row, "FECHA"), _field(row, "AÑO"), _field(row, "MES"))
            if period:
                trade.append({
                    "period": period,
                    "direction": direction,
                    "product": str(_field(row, "PRODUCTO") or "").strip(),
                    "volumeMbbl": _number(_field(row, "MBLS")),
                })
    trade_monthly = _group_sum(trade, ("period", "direction", "product"), "volumeMbbl")

    structure = []
    structure_path = _workbook_path(root, "MNHL3-")
    for row in _records(structure_path, "Estructura de precios", header_row=5):
        period = _period(_field(row, "Fecha"))
        price = _number(_field(row, "Precio venta público"))
        if period and price > 0:
            structure.append({
                "period": period,
                "product": str(_field(row, "COMBUSTIBLES") or "").strip(),
                "publicPrice": price,
            })

    return {
        "fines": fine_counts,
        "fuelPrices": monthly_prices,
        "fuelSales": _group_sum(sales, ("period", "product"), "gallons"),
        "glpPrices": glp_average_prices,
        "glpSales": _group_sum(glp_sales, ("period", "product"), "barrels"),
        "trade": trade_monthly,
        "priceStructure": structure,
    }


def build_operational_data(imports_path: Path) -> dict[str, Any]:
    return {
        "naturalGas": _read_gas(imports_path),
        "liquidFuels": _read_fuels(imports_path),
    }
